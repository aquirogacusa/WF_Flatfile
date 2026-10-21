import glob
import os
import shutil
import sys
from datetime import datetime

import openpyxl
import pandas as pd
from dotenv import load_dotenv

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

KEY_FILE_DATE_FMT = "%m.%d.%Y"

# Columna de upload -> (alias del encabezado en maestro, posicion de respaldo 1-based)
MAESTRO_SOURCES = {
    "Customer Code": (["customer code"], 1),
    "Customer Name": (["name 1"], 6),
    "Street": (["street"], 8),
    "City": (["city"], 9),
    "Region": (["region"], 10),
    "Zip Code": (["postal code"], 12),
    "Sales Rep Code": (["number"], 38),
    "Sales Rep Name": (["sales representative name surname"], 39),
    "Route Code": (["partner description"], 40),
    "Territory Code": (["group of sallers", "group of sellers"], 41),
    "Territory Name": (["sales group description"], 42),
    "Team Leader Code": (["z5 code"], 47),
    "Team Leader Name": (["z5 name"], 48),
    "Store Type (ABC)": (["abc class.", "abc class"], 29),
}
FILTER_ALIAS = "number"
FILTER_FALLBACK_POS = 38  # Columna AL (Number) en Maestro Clientes


def excel_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def code_variants(text):
    variants = {text}
    if text.isdigit():
        variants.add(text.lstrip("0"))
    return variants


def load_rep_codes(path):
    raw = pd.read_excel(path, header=None, dtype=object, usecols=[0])
    codes = set()
    for value in raw.iloc[:, 0]:
        text = excel_text(value)
        if text and any(ch.isdigit() for ch in text):
            codes |= code_variants(text)
    if not codes:
        raise ValueError(f"No se encontro ningun codigo de vendedor en {path}")
    return codes


def resolve_maestro_path(env):
    source_folder = env.get("SOURCE_FOLDER", "").strip()
    if not source_folder or not os.path.isdir(source_folder):
        raise FileNotFoundError(f"No existe la carpeta SOURCE_FOLDER: {source_folder}")
    date = datetime.now().strftime(KEY_FILE_DATE_FMT)
    pattern = env.get("MAESTRO_FILE_NAME", "Maestro Clientes (Pre) {date}.xlsx").strip()
    exact = os.path.join(source_folder, pattern.replace("{date}", date))
    if os.path.isfile(exact):
        return exact
    stem = os.path.splitext(pattern)[0].replace("{date}", "")
    candidates = sorted(
        glob.glob(os.path.join(source_folder, f"{stem}*.xlsx")),
        key=os.path.getmtime,
        reverse=True,
    )
    if not candidates:
        candidates = sorted(
            glob.glob(os.path.join(source_folder, "Maestro Clientes*.xlsx")),
            key=os.path.getmtime,
            reverse=True,
        )
    for candidate in candidates:
        if not os.path.basename(candidate).startswith("~$"):
            print(f"No existe el archivo del dia; se usara el mas reciente: {candidate}")
            return candidate
    raise FileNotFoundError(f"No se encontro el archivo Maestro Clientes en {source_folder}")


def build_output_path(env):
    date = datetime.now().strftime(KEY_FILE_DATE_FMT)
    file_name = env.get("UPLOAD_FILE_NAME", "Upload EasyRoute {date}.xlsx").strip()
    return os.path.join(env["SOURCE_FOLDER"], file_name.replace("{date}", date))


def resolve_project_file(env, key, default_name):
    path = env.get(key, default_name).strip() or default_name
    return path if os.path.isabs(path) else os.path.join(BASE_DIR, path)


def detect_column_positions(header_row):
    norm = [excel_text(v).lower() for v in header_row]
    positions = {}
    for upload_name, (aliases, fallback) in MAESTRO_SOURCES.items():
        pos = next((norm.index(a) + 1 for a in aliases if a in norm), None)
        positions[upload_name] = pos if pos else fallback
    filter_pos = norm.index(FILTER_ALIAS) + 1 if FILTER_ALIAS in norm else FILTER_FALLBACK_POS
    return positions, filter_pos


def run(env):
    maestro_path = resolve_maestro_path(env)
    output_path = build_output_path(env)

    reps_path = resolve_project_file(env, "EASYROUTE_REPS_FILE", "Sales Rep DSD.xlsx")
    template_path = resolve_project_file(env, "EASYROUTE_TEMPLATE_FILE", "BaseCustomersDSD_Empty.xlsx")
    for p in (reps_path, template_path):
        if not os.path.isfile(p):
            raise FileNotFoundError(f"No se encontró el archivo: {p}")

    rep_codes = load_rep_codes(reps_path)
    print(f"Filtro DSD: {len(rep_codes)} codigos de vendedor cargados desde {reps_path}")
    print(f"Leyendo maestro: {maestro_path}")

    wb_src = openpyxl.load_workbook(maestro_path, read_only=True)
    ws_src = wb_src.active
    rows = ws_src.iter_rows(values_only=True)
    header_row = next(rows, None)
    if header_row is None:
        raise ValueError(f"El maestro esta vacio: {maestro_path}")
    positions, filter_pos = detect_column_positions(header_row)

    shutil.copy2(template_path, output_path)
    wb_out = openpyxl.load_workbook(output_path)
    ws_out = wb_out.active
    upload_headers = [excel_text(ws_out.cell(1, c).value) for c in range(1, ws_out.max_column + 1)]

    total = 0
    matched = 0
    out_row = 2
    for row in rows:
        if all(excel_text(v) == "" for v in row):
            continue
        total += 1
        key = excel_text(row[filter_pos - 1]) if len(row) >= filter_pos else ""
        if not key or not (code_variants(key) & rep_codes):
            continue
        matched += 1
        for col, name in enumerate(upload_headers, start=1):
            source = MAESTRO_SOURCES.get(name)
            if not source:
                continue
            pos = positions[name]
            value = row[pos - 1] if len(row) >= pos else None
            if value is not None:
                ws_out.cell(out_row, col, value)
        out_row += 1

    wb_src.close()
    wb_out.save(output_path)
    print(f"Filas procesadas: {matched} de {total} filas del maestro")
    print(f"Archivo generado: {output_path}")
    return output_path


def main():
    load_dotenv()
    env = {k: v for k, v in os.environ.items() if v is not None}
    try:
        run(env)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
