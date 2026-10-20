import os
import shutil
import sys
from copy import copy
from datetime import datetime

import openpyxl
import pandas as pd
from dotenv import load_dotenv
from openpyxl.utils import get_column_letter

KEY_COL_IDX = 2          # Columna C (Customer) en ambos archivos
Z5_CODE_COL_IDX = 36     # Columna AK (Number) en Customers_SAP_Z5.xlsx
Z5_NAME_COL_IDX = 37     # Columna AL (Sales Representative Name Surname)


def excel_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_z5(path):
    raw = pd.read_excel(path, header=None, dtype=object)
    first = raw.iloc[0]
    key_txt = excel_text(first[KEY_COL_IDX]).lower() if len(first) > KEY_COL_IDX else ""
    if key_txt == "customer" or all(excel_text(v) == "" for v in first):
        raw = raw.iloc[1:]
    return raw


def build_output_path(env):
    date = datetime.now().strftime("%m.%d.%Y")
    full_name = env.get("OUTPUT_FILE_NAME", "").strip()
    if full_name:
        file_name = full_name.replace("{date}", date)
    else:
        base_name = env.get("OUTPUT_BASE_NAME", "Maestro Clientes").strip()
        file_name = f"{base_name} {date}.xlsx"
    return os.path.join(env["SOURCE_FOLDER"], file_name)


def run(env):
    source_folder = env.get("SOURCE_FOLDER", "").strip()
    base_file = env.get("BASE_FILE_NAME", "Customers_SAP.xlsx").strip()
    z5_file = env.get("Z5_FILE_NAME", "Customers_SAP_Z5.xlsx").strip()

    if not source_folder or not os.path.isdir(source_folder):
        raise FileNotFoundError(f"No existe la carpeta SOURCE_FOLDER: {source_folder}")

    base_path = os.path.join(source_folder, base_file)
    z5_path = os.path.join(source_folder, z5_file)
    for p in (base_path, z5_path):
        if not os.path.isfile(p):
            raise FileNotFoundError(f"No se encontró el archivo: {p}")

    print(f"Leyendo Z5: {z5_path}")
    z5_df = load_z5(z5_path)
    for label, idx in (("Customer", KEY_COL_IDX), ("Number", Z5_CODE_COL_IDX), ("Sales Rep Name", Z5_NAME_COL_IDX)):
        if idx >= z5_df.shape[1]:
            raise IndexError(f"{z5_file}: se esperaba la columna '{label}' en la posicion {idx + 1} pero el archivo solo tiene {z5_df.shape[1]} columnas")

    z5_lookup = {}
    for _, row in z5_df.iterrows():
        key = excel_text(row[KEY_COL_IDX])
        if key and key not in z5_lookup:
            z5_lookup[key] = (row[Z5_CODE_COL_IDX], row[Z5_NAME_COL_IDX])

    print(f"Leyendo base (con formato): {base_path}")
    output_path = build_output_path(env)
    shutil.copy2(base_path, output_path)
    wb = openpyxl.load_workbook(output_path)
    ws = wb.active

    ncols = ws.max_column
    nrows = ws.max_row
    for expected, idx in (("Customer", KEY_COL_IDX), ("Number", Z5_CODE_COL_IDX), ("Sales Rep", Z5_NAME_COL_IDX)):
        if idx >= ncols:
            raise IndexError(f"{base_file}: se esperaba una columna en la posicion {idx + 1} pero solo tiene {ncols}")

    # Fila de encabezado: si la fila 1 no existe o esta vacia, se reserva para los nombres
    c1 = excel_text(ws.cell(row=1, column=KEY_COL_IDX + 1).value).lower()
    if c1 == "customer":
        data_start = 2
    elif all(excel_text(ws.cell(row=1, column=c).value) == "" for c in range(1, ncols + 1)):
        data_start = 2
    else:
        ws.insert_rows(1)
        nrows += 1
        data_start = 2

    ws.insert_cols(1)
    ncols += 1
    cust_col = KEY_COL_IDX + 2
    code_col, name_col = ncols + 1, ncols + 2

    hdr_style = ws.cell(1, 2)._style
    for col, name in ((1, "Customer Code"), (code_col, "Z5 CODE"), (name_col, "Z5 NAME")):
        cell = ws.cell(1, col, name)
        cell._style = copy(hdr_style)

    matched = 0
    for r in range(data_start, nrows + 1):
        key = excel_text(ws.cell(r, cust_col).value)
        base_style = ws.cell(r, cust_col)._style
        ws.cell(r, 1, key[-8:] if key else None)._style = copy(base_style)
        if key in z5_lookup:
            matched += 1
            code_val, name_val = z5_lookup[key]
            ws.cell(r, code_col, code_val if code_val is not None else None)._style = copy(base_style)
            ws.cell(r, name_col, name_val if name_val is not None else None)._style = copy(base_style)

    for col, width in ((1, 14), (code_col, 12), (name_col, 35)):
        ws.column_dimensions[get_column_letter(col)].width = width

    print(f"Coincidencias Z5: {matched} de {nrows - data_start + 1} filas")
    wb.save(output_path)
    print(f"Archivo generado: {output_path}")
    return output_path


def main():
    load_dotenv()
    env = {k: v for k, v in os.environ.items() if v is not None}
    try:
        run(env)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
